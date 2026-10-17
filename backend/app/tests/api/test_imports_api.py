"""API tests for spreadsheet import (``docs/modules/17-data-import-export.md``).

Real app, services, repositories and database. Records are created through each
module's own service, so these tests also prove MRNs and uniqueness work as
they do by hand.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.api.dependencies.db import get_db_session
from app.api.dependencies.services import get_audit_sink
from app.main import create_app
from app.models.billing import Service
from app.models.lab import LabTest
from app.models.patient import Patient
from app.models.pharmacy import Medicine
from app.tests.billing_helpers import auth_headers, insert_patient, insert_user_with_permissions
from app.tests.conftest import RecordingAuditSink
from app.tests.pharmacy_helpers import insert_medicine
from app.utils.spreadsheet import XLSX_TYPE

if TYPE_CHECKING:
    import uuid
    from collections.abc import AsyncGenerator

    from fastapi import FastAPI
    from sqlalchemy.ext.asyncio import AsyncSession

pytestmark = pytest.mark.database

URL = "/api/v1/imports"

PATIENT_HEADERS = [
    "First name",
    "Last name",
    "Date of birth (YYYY-MM-DD)",
    "Gender",
    "Phone",
    "Address line 1",
    "City",
]


def _xlsx(headers: list[str], rows: list[list[Any]]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def audit() -> RecordingAuditSink:
    return RecordingAuditSink()


@pytest_asyncio.fixture
async def api(db_session: AsyncSession, audit: RecordingAuditSink) -> AsyncGenerator[AsyncClient]:
    application: FastAPI = create_app()

    async def _session_override() -> AsyncGenerator[AsyncSession]:
        yield db_session

    application.dependency_overrides[get_db_session] = _session_override
    application.dependency_overrides[get_audit_sink] = lambda: audit
    transport = ASGITransport(app=application)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        yield client
    application.dependency_overrides.clear()


@pytest_asyncio.fixture
async def admin(db_session: AsyncSession, hospital_id: uuid.UUID) -> dict[str, str]:
    user = await insert_user_with_permissions(db_session, hospital_id, ["data.import"])
    return auth_headers(user.id, hospital_id)


async def _upload(
    api: AsyncClient,
    headers: dict[str, str],
    kind: str,
    data: bytes,
    *,
    dry_run: bool,
    filename: str = "upload.xlsx",
    content_type: str = XLSX_TYPE,
) -> Any:
    return await api.post(
        f"{URL}/{kind}",
        params={"dry_run": str(dry_run).lower()},
        content=data,
        headers={**headers, "Content-Type": content_type, "X-Filename": filename},
    )


async def _count(session: AsyncSession, model: Any, hospital_id: uuid.UUID) -> int:
    stmt = select(func.count()).select_from(model).where(model.hospital_id == hospital_id)
    return int((await session.execute(stmt)).scalar_one())


class TestTemplates:
    @pytest.mark.parametrize("kind", ["patients", "services", "medicines", "lab-tests"])
    async def test_each_kind_has_a_template_with_instructions(
        self, api: AsyncClient, admin: dict[str, str], kind: str
    ) -> None:
        response = await api.get(f"{URL}/{kind}/template", headers=admin)

        assert response.status_code == 200
        assert response.headers["content-type"] == XLSX_TYPE
        assert f"atheris-{kind}-template.xlsx" in response.headers["content-disposition"]
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames[1] == "Instructions"

    async def test_an_unknown_kind_is_not_found(
        self, api: AsyncClient, admin: dict[str, str]
    ) -> None:
        response = await api.get(f"{URL}/doctors/template", headers=admin)

        assert response.status_code == 404


class TestPatients:
    ROWS = [
        ["Ananya", "Rao", "1990-01-31", "F", "9812345678", "12 MG Road", "Hyderabad"],
        ["Ravi", "Kumar", "15/08/1985", "male", "", "", ""],
    ]

    async def test_a_dry_run_reports_and_saves_nothing(
        self,
        api: AsyncClient,
        admin: dict[str, str],
        db_session: AsyncSession,
        hospital_id: uuid.UUID,
        audit: RecordingAuditSink,
    ) -> None:
        response = await _upload(
            api, admin, "patients", _xlsx(PATIENT_HEADERS, self.ROWS), dry_run=True
        )

        assert response.status_code == 200, response.text
        report = response.json()["data"]
        assert report["dry_run"] is True
        assert (report["total_rows"], report["ready"], report["created"]) == (2, 2, 0)
        assert report["problems"] == []
        assert await _count(db_session, Patient, hospital_id) == 0
        assert audit.events == []

    async def test_a_real_run_registers_patients_with_mrns(
        self,
        api: AsyncClient,
        admin: dict[str, str],
        db_session: AsyncSession,
        hospital_id: uuid.UUID,
        audit: RecordingAuditSink,
    ) -> None:
        response = await _upload(
            api, admin, "patients", _xlsx(PATIENT_HEADERS, self.ROWS), dry_run=False
        )

        report = response.json()["data"]
        assert report["created"] == 2
        patients = (
            (await db_session.execute(select(Patient).where(Patient.hospital_id == hospital_id)))
            .scalars()
            .all()
        )
        by_name = {p.first_name: p for p in patients}
        assert by_name["Ananya"].phone == "+919812345678"
        assert by_name["Ravi"].date_of_birth.isoformat() == "1985-08-15"
        assert all(p.mrn for p in patients)
        assert "data.import.completed" in [e.action for e in audit.events]

    async def test_someone_already_registered_is_skipped_not_overwritten(
        self,
        api: AsyncClient,
        admin: dict[str, str],
        db_session: AsyncSession,
        hospital_id: uuid.UUID,
    ) -> None:
        # insert_patient makes "Ananya Rao", born 1990-01-01.
        existing = await insert_patient(db_session, hospital_id)
        rows = [
            ["ananya", "RAO", "1990-01-01", "female", "", "", ""],
            ["Meera", "Shah", "2001-03-04", "female", "", "", ""],
            ["Meera", "Shah", "2001-03-04", "female", "", "", ""],
        ]

        response = await _upload(
            api, admin, "patients", _xlsx(PATIENT_HEADERS, rows), dry_run=False
        )

        report = response.json()["data"]
        assert report["created"] == 1
        skipped = {s["row"]: s["reason"] for s in report["skipped"]}
        assert existing.mrn in skipped[2]
        assert "earlier row" in skipped[4]
        assert await _count(db_session, Patient, hospital_id) == 2

    async def test_bad_rows_are_reported_by_row_and_column_and_good_ones_still_import(
        self,
        api: AsyncClient,
        admin: dict[str, str],
        db_session: AsyncSession,
        hospital_id: uuid.UUID,
    ) -> None:
        rows = [
            ["Ananya", "Rao", "31 Jan 1990", "female", "", "", ""],
            ["", "Kumar", "1985-08-15", "male", "", "", ""],
            ["Ravi", "Kumar", "1985-08-15", "robot", "", "", ""],
            ["Meera", "Shah", "2001-03-04", "female", "12", "", ""],
            ["Arjun", "Das", "1979-11-02", "M", "", "", ""],
        ]

        response = await _upload(
            api, admin, "patients", _xlsx(PATIENT_HEADERS, rows), dry_run=False
        )

        report = response.json()["data"]
        problems = {(p["row"], p["column"]) for p in report["problems"]}
        assert (2, "Date of birth") in problems
        assert (3, "First name") in problems
        assert (4, "Gender") in problems
        assert (5, "Phone") in problems
        assert report["created"] == 1
        assert await _count(db_session, Patient, hospital_id) == 1

    async def test_a_city_without_a_street_is_reported_clearly(
        self, api: AsyncClient, admin: dict[str, str]
    ) -> None:
        rows = [["Ananya", "Rao", "1990-01-31", "female", "", "", "Hyderabad"]]

        response = await _upload(api, admin, "patients", _xlsx(PATIENT_HEADERS, rows), dry_run=True)

        [problem] = response.json()["data"]["problems"]
        assert problem["column"] == "Address line 1"
        assert "leave the address empty" in problem["message"]

    async def test_a_missing_required_column_is_reported_before_any_row(
        self, api: AsyncClient, admin: dict[str, str]
    ) -> None:
        data = _xlsx(["First name", "Last name", "Favourite colour"], [["Ananya", "Rao", "blue"]])

        response = await _upload(api, admin, "patients", data, dry_run=True)

        report = response.json()["data"]
        assert {p["column"] for p in report["problems"]} == {"Date of birth", "Gender"}
        assert report["ready"] == 0

    async def test_previous_record_number_is_kept_in_the_notes(
        self,
        api: AsyncClient,
        admin: dict[str, str],
        db_session: AsyncSession,
        hospital_id: uuid.UUID,
    ) -> None:
        data = _xlsx(
            ["First name", "Last name", "Date of birth", "Gender", "Previous record no.", "Extra"],
            [["Ananya", "Rao", "1990-01-31", "female", "OLD-1042", "x"]],
        )

        response = await _upload(api, admin, "patients", data, dry_run=False)

        assert response.json()["data"]["ignored_columns"] == ["extra"]
        patient = (
            await db_session.execute(select(Patient).where(Patient.hospital_id == hospital_id))
        ).scalar_one()
        assert "OLD-1042" in (patient.notes or "")


class TestCatalogs:
    async def test_services_from_csv(
        self,
        api: AsyncClient,
        admin: dict[str, str],
        db_session: AsyncSession,
        hospital_id: uuid.UUID,
    ) -> None:
        data = b'Code,Name,Category,Price,Taxable\ncons-gen,General consultation,Consultation,500,no\nXR,X-ray,Radiology,"1,200.00",yes\n'

        response = await _upload(
            api, admin, "services", data, dry_run=False, filename="s.csv", content_type="text/csv"
        )

        assert response.json()["data"]["created"] == 2, response.text
        codes = (
            (
                await db_session.execute(
                    select(Service.code).where(Service.hospital_id == hospital_id)
                )
            )
            .scalars()
            .all()
        )
        assert sorted(codes) == ["CONS-GEN", "XR"]

    async def test_a_medicine_whose_sku_exists_is_skipped(
        self,
        api: AsyncClient,
        admin: dict[str, str],
        db_session: AsyncSession,
        hospital_id: uuid.UUID,
    ) -> None:
        await insert_medicine(db_session, hospital_id, "PARA-500")
        data = _xlsx(
            ["SKU", "Name", "Price", "Prescription required"],
            [["para-500", "Paracetamol", "2.5", "no"], ["AMOX-250", "Amoxicillin", "8", "yes"]],
        )

        response = await _upload(api, admin, "medicines", data, dry_run=False)

        report = response.json()["data"]
        assert report["created"] == 1
        assert report["skipped"][0]["row"] == 2
        assert await _count(db_session, Medicine, hospital_id) == 2

    async def test_lab_tests_with_a_reference_range(
        self,
        api: AsyncClient,
        admin: dict[str, str],
        db_session: AsyncSession,
        hospital_id: uuid.UUID,
    ) -> None:
        data = _xlsx(
            ["Code", "Name", "Unit", "Result type", "Reference low", "Reference high", "Price"],
            [["HB", "Haemoglobin", "g/dL", "numeric", 12, 15.5, 250]],
        )

        response = await _upload(api, admin, "lab-tests", data, dry_run=False)

        assert response.json()["data"]["created"] == 1, response.text
        test = (
            await db_session.execute(select(LabTest).where(LabTest.hospital_id == hospital_id))
        ).scalar_one()
        assert test.reference_ranges[0]["low"] in (12, "12", "12.0000", 12.0)


class TestRefusals:
    async def test_a_file_that_is_not_a_spreadsheet_is_refused(
        self, api: AsyncClient, admin: dict[str, str]
    ) -> None:
        response = await _upload(
            api,
            admin,
            "patients",
            b"%PDF-1.7",
            dry_run=True,
            filename="x.pdf",
            content_type="application/pdf",
        )

        assert response.status_code == 422
        assert "Excel" in response.json()["message"]

    async def test_a_declared_size_over_the_limit_is_refused_before_reading(
        self, api: AsyncClient, admin: dict[str, str]
    ) -> None:
        response = await api.post(
            f"{URL}/patients",
            content=b"x",
            headers={
                **admin,
                "Content-Type": "text/csv",
                "X-Filename": "x.csv",
                "Content-Length": str(6 * 1024 * 1024),
            },
        )

        assert response.status_code == 422

    async def test_without_the_permission_it_is_forbidden(
        self, api: AsyncClient, db_session: AsyncSession, hospital_id: uuid.UUID
    ) -> None:
        user = await insert_user_with_permissions(db_session, hospital_id, ["patient.create"])
        headers = auth_headers(user.id, hospital_id)

        assert (await api.get(f"{URL}/patients/template", headers=headers)).status_code == 403
        response = await _upload(api, headers, "patients", _xlsx(PATIENT_HEADERS, []), dry_run=True)
        assert response.status_code == 403

    async def test_without_a_token_it_is_unauthorised(self, api: AsyncClient) -> None:
        assert (await api.get(f"{URL}/patients/template")).status_code == 401

    async def test_patients_land_only_in_the_importers_hospital(
        self,
        api: AsyncClient,
        db_session: AsyncSession,
        hospital_id: uuid.UUID,
        other_hospital_id: uuid.UUID,
    ) -> None:
        other = await insert_user_with_permissions(db_session, other_hospital_id, ["data.import"])
        # The same person is already registered in the first hospital ...
        await insert_patient(db_session, hospital_id)
        data = _xlsx(PATIENT_HEADERS, [["Ananya", "Rao", "1990-01-01", "female", "", "", ""]])

        response = await _upload(
            api, auth_headers(other.id, other_hospital_id), "patients", data, dry_run=False
        )

        # ... which the second hospital neither sees nor counts as a duplicate.
        assert response.json()["data"]["created"] == 1
        assert await _count(db_session, Patient, other_hospital_id) == 1
        assert await _count(db_session, Patient, hospital_id) == 1
